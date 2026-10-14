"""Excel export. Builds the workbook in memory (fine up to EXPORT_MAX_ROWS);
larger exports should call the same function from a Celery task and upload
the bytes to MinIO instead of returning them directly."""
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _excel_value(col, value):
    if value is None:
        return None
    if col.type == "bool":
        return "Yes" if value else "No"
    if col.type == "datetime" and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)   # Excel has no tz support
    return value


def build_workbook(*, title, columns, rows, business=None, include_serial=True, generated_at=None, streaming=False):
    """columns: registry Col objects; rows: iterable of tuples aligned with columns.

    streaming=False: full formatting (merged banner, borders, auto column widths).
    streaming=True : openpyxl write-only mode for very large exports. Rows are written
                     as they arrive, so memory stays flat; the banner is centred across
                     the columns instead of merged and widths come from the headings."""
    generated_at = generated_at or timezone.localtime()
    builder = _build_streaming if streaming else _build_full
    return builder(title=title, columns=columns, rows=rows, business=business,
                   include_serial=include_serial, generated_at=generated_at)


def _build_full(*, title, columns, rows, business, include_serial, generated_at):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    offset = 1 if include_serial else 0
    ncols = len(columns) + offset
    r = 1

    def banner(text, font):
        nonlocal r
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
        if ncols > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        r += 1

    if business:
        banner(business["name"], Font(bold=True, size=14))
        for line in (business.get("address"), business.get("contact")):
            if line:
                banner(line, Font(size=10))
        r += 1
    banner(f"{title} Report", Font(bold=True, size=12))
    banner(f"Generated on: {generated_at:%d-%b-%Y %H:%M}", Font(italic=True, size=9))
    r += 1

    # header row
    header_row = r
    heads = (["S.No"] if include_serial else []) + [c.label for c in columns]
    widths = [len(h) for h in heads]
    for ci, head in enumerate(heads, start=1):
        cell = ws.cell(row=r, column=ci, value=head)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1

    # data rows
    n = 0
    for n, row in enumerate(rows, start=1):
        values = ([n] if include_serial else []) + [_excel_value(c, v) for c, v in zip(columns, row)]
        for ci, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=ci, value=val)
            if isinstance(val, str) and val.startswith("="):
                cell.data_type = "s"          # user data must never turn into a formula
            cell.border = BORDER
            if ci > offset:
                ctype = columns[ci - offset - 1].type
                if ctype == "datetime":
                    cell.number_format = "dd-mmm-yyyy hh:mm"
                elif ctype == "date":
                    cell.number_format = "dd-mmm-yyyy"
            if val is not None:
                widths[ci - 1] = max(widths[ci - 1], len(str(val)))
        r += 1

    last_data_row = r - 1
    if n:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{last_data_row}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.cell(row=r + 1, column=1, value=f"Total records: {n}").font = Font(bold=True)

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 50)

    # print setup: landscape, one page wide, header row repeats on every page
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"{header_row}:{header_row}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


_NUMBER_FORMATS = {"datetime": "dd-mmm-yyyy hh:mm", "date": "dd-mmm-yyyy"}


def _stream_cell(ws, col, value):
    value = _excel_value(col, value)
    formula_like = isinstance(value, str) and value.startswith("=")
    fmt = _NUMBER_FORMATS.get(col.type) if value is not None else None
    if not formula_like and not fmt:
        return value                                   # plain value: fastest path
    cell = WriteOnlyCell(ws, value=value)
    if formula_like:
        cell.data_type = "s"                           # user data must never turn into a formula
    if fmt:
        cell.number_format = fmt
    return cell


def _build_streaming(*, title, columns, rows, business, include_serial, generated_at):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title[:31])
    offset = 1 if include_serial else 0
    ncols = len(columns) + offset
    heads = (["S.No"] if include_serial else []) + [c.label for c in columns]

    # the banner lines above the table: (text, font)
    prelude = []
    if business:
        prelude.append((business["name"], Font(bold=True, size=14)))
        for line in (business.get("address"), business.get("contact")):
            if line:
                prelude.append((line, Font(size=10)))
        prelude.append((None, None))
    prelude.append((f"{title} Report", Font(bold=True, size=12)))
    prelude.append((f"Generated on: {generated_at:%d-%b-%Y %H:%M}", Font(italic=True, size=9)))
    prelude.append((None, None))
    header_row = len(prelude) + 1

    # sheet-level settings must be set before the first row is written
    ws.freeze_panes = f"A{header_row + 1}"
    for i, head in enumerate(heads, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(head) + 4, 12), 40)

    centre = Alignment(horizontal="centerContinuous")
    for text, font in prelude:
        if text is None:
            ws.append([])
            continue
        cells = []
        for i in range(ncols):
            cell = WriteOnlyCell(ws, value=text if i == 0 else None)
            cell.alignment = centre
            if i == 0:
                cell.font = font
            cells.append(cell)
        ws.append(cells)

    header_cells = []
    for head in heads:
        cell = WriteOnlyCell(ws, value=head)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_cells.append(cell)
    ws.append(header_cells)

    n = 0
    for n, row in enumerate(rows, start=1):
        values = [_stream_cell(ws, c, v) for c, v in zip(columns, row)]
        ws.append(([n] if include_serial else []) + values)

    if n:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{header_row + n}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.append([])
    total = WriteOnlyCell(ws, value=f"Total records: {n}")
    total.font = Font(bold=True)
    ws.append([total])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

PARENT_FILL = PatternFill("solid", fgColor="2E75B6")

def _head_fill(columns, i):
    return PARENT_FILL if 0 <= i < len(columns) and columns[i].group else HEADER_FILL